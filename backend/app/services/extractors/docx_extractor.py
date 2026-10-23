"""DOCX text extractor."""

import logging
from docx import Document

from app.services.extractors.base import ExtractionResult

logger = logging.getLogger(__name__)


def extract_docx(file_path: str) -> ExtractionResult:
    """Extract text from a .docx file.

    Extracts paragraphs as well as text inside tables.
    """
    doc = Document(file_path)
    paragraphs = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            paragraphs.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs.append(" | ".join(row_text))

    full_text = "\n\n".join(paragraphs)
    # Estimate page count roughly (~3000 chars or 500 words per page, minimum 1)
    word_count = len(full_text.split())
    page_count = max(1, (word_count + 499) // 500) if full_text else 1

    return ExtractionResult(
        text=full_text,
        page_count=page_count,
        method="docx",
    )
